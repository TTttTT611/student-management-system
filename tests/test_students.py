from io import BytesIO

from openpyxl import Workbook, load_workbook


def test_crud(client, admin, student_payload):
    res = client.post("/api/students", json=student_payload, headers=admin)
    assert res.status_code == 201, res.text
    sid = res.json()["id"]
    assert res.json()["class_name"] is None

    res = client.get(f"/api/students/{sid}", headers=admin)
    assert res.json()["name"] == "Alice Smith"

    res = client.put(f"/api/students/{sid}", json={"age": 21}, headers=admin)
    assert res.status_code == 200 and res.json()["age"] == 21

    assert client.delete(f"/api/students/{sid}", headers=admin).status_code == 204
    assert client.get(f"/api/students/{sid}", headers=admin).status_code == 404


def test_duplicate_student_no(client, admin, student_payload):
    assert client.post("/api/students", json=student_payload, headers=admin).status_code == 201
    res = client.post("/api/students", json=student_payload, headers=admin)
    assert res.status_code == 400 and "Student number" in res.json()["detail"]

    other = {**student_payload, "student_no": "S002"}
    sid = client.post("/api/students", json=other, headers=admin).json()["id"]
    res = client.put(f"/api/students/{sid}", json={"student_no": "S001"}, headers=admin)
    assert res.status_code == 400


def test_validation(client, admin, student_payload):
    for bad in (
        {"age": 0},
        {"age": 200},
        {"gender": "other"},
        {"name": ""},
        {"student_no": "has space"},
    ):
        res = client.post("/api/students", json={**student_payload, **bad}, headers=admin)
        assert res.status_code == 422, bad

    res = client.post("/api/students", json={**student_payload, "class_id": 999}, headers=admin)
    assert res.status_code == 400


def test_pagination_and_search(client, admin, student_payload):
    for i in range(25):
        payload = {**student_payload, "student_no": f"S{i:03d}", "name": f"Student {i}"}
        if i % 5 == 0:
            payload["major"] = "Mathematics"
        assert client.post("/api/students", json=payload, headers=admin).status_code == 201

    res = client.get("/api/students", params={"page": 1, "size": 10}, headers=admin).json()
    assert res["total"] == 25 and len(res["items"]) == 10
    res = client.get("/api/students", params={"page": 3, "size": 10}, headers=admin).json()
    assert len(res["items"]) == 5

    res = client.get("/api/students", params={"keyword": "Mathematics"}, headers=admin).json()
    assert res["total"] == 5
    res = client.get("/api/students", params={"keyword": "S007"}, headers=admin).json()
    assert res["total"] == 1 and res["items"][0]["name"] == "Student 7"


def test_class_assignment(client, admin, student_payload):
    cid = client.post("/api/classes", json={"name": "CS-1", "major": "Computer Science"}, headers=admin).json()["id"]
    res = client.post("/api/students", json={**student_payload, "class_id": cid}, headers=admin)
    assert res.json()["class_name"] == "CS-1"

    res = client.get("/api/students", params={"class_id": cid}, headers=admin).json()
    assert res["total"] == 1
    assert client.get(f"/api/classes/{cid}", headers=admin).json()["student_count"] == 1

    # deleting a class keeps its students, only detaches them
    assert client.delete(f"/api/classes/{cid}", headers=admin).status_code == 204
    res = client.get("/api/students", headers=admin).json()
    assert res["total"] == 1 and res["items"][0]["class_id"] is None


def test_export_import(client, admin, student_payload):
    client.post("/api/classes", json={"name": "Class A"}, headers=admin)
    client.post("/api/students", json=student_payload, headers=admin)

    res = client.get("/api/students/export", headers=admin)
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0] == ("Student No", "Name", "Age", "Gender", "Major", "Class")
    assert rows[1][:2] == ("S001", "Alice Smith")

    wb = Workbook()
    ws = wb.active
    ws.append(["Student No", "Name", "Age", "Gender", "Major", "Class"])
    ws.append(["S001", "Alice Smith", 20, "female", "Computer Science", None])  # exists -> skipped
    ws.append(["S002", "Bob Lee", 19, "male", "Mathematics", "Class A"])  # ok
    ws.append(["S003", "Carol Wang", 19, "female", None, "No Such Class"])  # unknown class
    ws.append(["S004", "Dan Zhao", 999, "male", None, None])  # invalid age
    buf = BytesIO()
    wb.save(buf)

    res = client.post(
        "/api/students/import",
        files={"file": ("students.xlsx", buf.getvalue(), "application/octet-stream")},
        headers=admin,
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["created"] == 1 and body["skipped"] == 1 and len(body["errors"]) == 2

    res = client.get("/api/students", params={"keyword": "Bob Lee"}, headers=admin).json()
    assert res["total"] == 1 and res["items"][0]["class_name"] == "Class A"

    res = client.post(
        "/api/students/import",
        files={"file": ("x.csv", b"a,b", "text/csv")},
        headers=admin,
    )
    assert res.status_code == 400
