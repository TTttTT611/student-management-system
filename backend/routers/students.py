from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import or_
from sqlalchemy.orm import Session, joinedload

from .. import models, schemas
from ..database import get_db
from ..security import get_current_user, require_admin

router = APIRouter(prefix="/students", tags=["students"], dependencies=[Depends(get_current_user)])

EXPORT_COLUMNS = ["Student No", "Name", "Age", "Gender", "Major", "Class"]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _to_out(s: models.Student) -> schemas.StudentOut:
    return schemas.StudentOut(
        id=s.id,
        name=s.name,
        age=s.age,
        gender=s.gender,
        student_no=s.student_no,
        major=s.major,
        class_id=s.class_id,
        class_name=s.clazz.name if s.clazz else None,
    )


def _get_or_404(db: Session, student_id: int) -> models.Student:
    student = (
        db.query(models.Student)
        .options(joinedload(models.Student.clazz))
        .filter(models.Student.id == student_id)
        .first()
    )
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")
    return student


def _check_class(db: Session, class_id: int | None) -> None:
    if class_id is not None and db.get(models.Clazz, class_id) is None:
        raise HTTPException(status_code=400, detail="Class not found")


def _check_student_no(db: Session, student_no: str, exclude_id: int | None = None) -> None:
    q = db.query(models.Student).filter(models.Student.student_no == student_no)
    if exclude_id is not None:
        q = q.filter(models.Student.id != exclude_id)
    if q.first():
        raise HTTPException(status_code=400, detail="Student number already exists")


@router.get("", response_model=schemas.Page[schemas.StudentOut])
def list_students(
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=100),
    keyword: str | None = Query(None, max_length=50),
    class_id: int | None = None,
    db: Session = Depends(get_db),
):
    q = db.query(models.Student).options(joinedload(models.Student.clazz))
    if keyword:
        like = f"%{keyword.strip()}%"
        q = q.filter(
            or_(
                models.Student.name.like(like),
                models.Student.student_no.like(like),
                models.Student.major.like(like),
            )
        )
    if class_id is not None:
        q = q.filter(models.Student.class_id == class_id)

    total = q.count()
    items = q.order_by(models.Student.id.desc()).offset((page - 1) * size).limit(size).all()
    return schemas.Page(items=[_to_out(s) for s in items], total=total, page=page, size=size)


@router.get("/export", dependencies=[Depends(require_admin)])
def export_students(db: Session = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(EXPORT_COLUMNS)
    students = (
        db.query(models.Student)
        .options(joinedload(models.Student.clazz))
        .order_by(models.Student.id)
        .all()
    )
    for s in students:
        ws.append([s.student_no, s.name, s.age, s.gender, s.major, s.clazz.name if s.clazz else None])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type=XLSX_MIME,
        headers={"Content-Disposition": 'attachment; filename="students.xlsx"'},
    )


@router.post("/import", dependencies=[Depends(require_admin)])
async def import_students(file: UploadFile, db: Session = Depends(get_db)):
    """Bulk import using the export layout: Student No, Name, Age, Gender, Major, Class. Rows whose student number already exists are skipped."""
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx files are supported")
    try:
        wb = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="File could not be parsed")

    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None or [str(h).strip() if h is not None else "" for h in header[:6]] != EXPORT_COLUMNS:
        raise HTTPException(status_code=400, detail=f"Header row must be: {', '.join(EXPORT_COLUMNS)}")

    classes_by_name = {c.name: c.id for c in db.query(models.Clazz).all()}
    existing_nos = {no for (no,) in db.query(models.Student.student_no).all()}

    created, skipped, errors = 0, 0, []
    for line_no, row in enumerate(rows, start=2):
        row = list(row) + [None] * (6 - len(row))
        student_no, name, age, gender, major, class_name = row[:6]
        if all(v is None or str(v).strip() == "" for v in row[:6]):
            continue
        student_no = str(student_no).strip() if student_no is not None else ""
        if student_no in existing_nos:
            skipped += 1
            continue
        class_id = None
        if class_name:
            class_id = classes_by_name.get(str(class_name).strip())
            if class_id is None:
                errors.append(f"Row {line_no}: class '{class_name}' does not exist")
                continue
        try:
            data = schemas.StudentCreate(
                student_no=student_no,
                name=str(name).strip() if name is not None else "",
                age=age,
                gender=str(gender).strip() if gender is not None else "",
                major=str(major).strip() if major else None,
                class_id=class_id,
            )
        except ValidationError as e:
            fields = ", ".join(str(err["loc"][0]) for err in e.errors())
            errors.append(f"Row {line_no}: invalid field(s): {fields}")
            continue
        db.add(models.Student(**data.model_dump()))
        existing_nos.add(student_no)
        created += 1

    db.commit()
    return {"created": created, "skipped": skipped, "errors": errors}


@router.get("/{student_id}", response_model=schemas.StudentOut)
def get_student(student_id: int, db: Session = Depends(get_db)):
    return _to_out(_get_or_404(db, student_id))


@router.post("", response_model=schemas.StudentOut, status_code=201, dependencies=[Depends(require_admin)])
def create_student(data: schemas.StudentCreate, db: Session = Depends(get_db)):
    _check_student_no(db, data.student_no)
    _check_class(db, data.class_id)
    student = models.Student(**data.model_dump())
    db.add(student)
    db.commit()
    db.refresh(student)
    return _to_out(student)


@router.put("/{student_id}", response_model=schemas.StudentOut, dependencies=[Depends(require_admin)])
def update_student(student_id: int, data: schemas.StudentUpdate, db: Session = Depends(get_db)):
    student = _get_or_404(db, student_id)
    changes = data.model_dump(exclude_unset=True)
    if "student_no" in changes and changes["student_no"] != student.student_no:
        _check_student_no(db, changes["student_no"], exclude_id=student_id)
    if "class_id" in changes:
        _check_class(db, changes["class_id"])
    for k, v in changes.items():
        setattr(student, k, v)
    db.commit()
    db.refresh(student)
    return _to_out(student)


@router.delete("/{student_id}", status_code=204, dependencies=[Depends(require_admin)])
def delete_student(student_id: int, db: Session = Depends(get_db)):
    student = _get_or_404(db, student_id)
    db.delete(student)
    db.commit()
